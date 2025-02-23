import re
from typing import List, Dict
import json
from langchain.schema import Document
import fitz  # PyMuPDF for PDF
from docx import Document  # python-docx for DOCX
from pptx import Presentation  # python-pptx for PPTX
import openpyxl  # openpyxl for XLSX
from dotenv import load_dotenv
import os

def extract_content(text):
    # Use regular expression to find content between [[ and ]]
    pattern = r'\[\[(.*?)\]\]'
    result = re.findall(pattern, text)
    return str(result)


def process_json_to_documents(data: List[Dict]) -> List[Dict]:
    if isinstance(data, bytes):
        data = data.decode('utf-8')
    
    if isinstance(data, str):
        data = json.loads(data)
    
    documents = []
    for item in data:
        # JSON 구조에 따라 적절히 수정
        text = f"{item[1]}"
        doc = Document(page_content=text, metadata={})
        documents.append(doc)
    return documents


def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    
    if ext == '.pdf':
        try:
            doc = fitz.open(file_path)
            for page in doc:
                text += page.get_text()
        except Exception as e:
            print(f"PDF 읽기 오류 ({file_path}): {e}")
            
    elif ext == '.docx':
        try:
            doc = Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs])
        except Exception as e:
            print(f"DOCX 읽기 오류 ({file_path}): {e}")
            
    elif ext == '.pptx':
        try:
            prs = Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        except Exception as e:
            print(f"PPTX 읽기 오류 ({file_path}): {e}")
            
    elif ext == '.xlsx':
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_text = " ".join([str(cell) for cell in row if cell is not None])
                    text += row_text + "\n"
        except Exception as e:
            print(f"XLSX 읽기 오류 ({file_path}): {e}")
            
    else:
        print(f"지원되지 않는 파일 형식: {ext}")
    
    return text