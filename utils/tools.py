import re
import json
from typing import List, Dict, Optional
from langchain.schema import Document
from langchain_core.pydantic_v1 import BaseModel, Field
from langchain_core.prompts import ChatPromptTemplate
from langchain_openai import ChatOpenAI
import fitz  # PyMuPDF
from docx import Document as DocxDocument
from pptx import Presentation
import openpyxl
import os

# --- Text Extraction from Documents ---

def extract_text(file_path: str) -> str:
    """Extracts text content from various file formats (PDF, DOCX, PPTX, XLSX)."""
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    try:
        if ext == '.pdf':
            with fitz.open(file_path) as doc:
                for page in doc:
                    text += page.get_text()
        elif ext == '.docx':
            doc = DocxDocument(file_path)
            text = "\n".join([para.text for para in doc.paragraphs])
        elif ext == '.pptx':
            prs = Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        elif ext == '.xlsx':
            wb = openpyxl.load_workbook(file_path, read_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_text = " ".join([str(cell) for cell in row if cell is not None])
                    text += row_text + "\n"
        else:
            print(f"Unsupported file format: {ext}")
    except Exception as e:
        print(f"Error reading {file_path}: {e}")
    return text

# --- Structured Data Extraction using Function Calling ---

class ExtractedEntity(BaseModel):
    """Represents a single extracted entity (e.g., a support program or financial product)."""
    name: str = Field(description="The name of the support program or financial product.")
    link: Optional[str] = Field(description="The URL link to the program or product, if available.")

class ExtractedData(BaseModel):
    """A list of all entities extracted from the text."""
    entities: List[ExtractedEntity]

def extract_structured_data(text: str, llm: ChatOpenAI) -> List[Dict]:
    """Extracts names and links of programs/products from text using LLM function calling."""
    if not text:
        return []

    prompt = ChatPromptTemplate.from_messages([
        ("system", """
        You are an expert at extracting specific information from text.
        Your task is to identify all mentions of government support programs or financial products and their corresponding URLs.
        Extract all of them. Do not extract general-purpose links like '전자공시시스템 (https://dart.fss.or.kr)'.
        If no relevant entities are found, return an empty list.
        """),
        ("human", "Please extract the entities from the following text:\n\n{input}")
    ])

    try:
        extractor_chain = prompt | llm.with_structured_output(ExtractedData)
        result = extractor_chain.invoke({"input": text})
        return [entity.dict() for entity in result.entities]
    except Exception as e:
        print(f"Error during structured data extraction: {e}")
        return []
